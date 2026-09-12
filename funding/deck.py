"""Three editable slides and matching PDF, generated from the frozen release."""
import json,math,textwrap
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.chart import XL_CHART_TYPE,XL_LEGEND_POSITION
from pptx.chart.data import CategoryChartData
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor
from pypdf import PdfReader

ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'outputs';OUT.mkdir(exist_ok=True)
r=json.loads((ROOT/'public/funding/current.json').read_text())
W,H=13.333,7.5
C={'paper':'F7F8F2','ink':'1B3E31','muted':'7A8C6D','line':'DCE3D2','green':'426D4D','sage':'E8EEDB','cream':'EBE4CD','white':'FFFEFA','dark':'173F34','SE':'3F8D77','US':'6F79B6','SG':'D38B4E'}
prs=Presentation();prs.slide_width=Inches(W);prs.slide_height=Inches(H)
prs.core_properties.title='Longview | Ageing research funding';prs.core_properties.author='Selja, Max, Jan';prs.core_properties.subject='Stockholm AI × Longevity Hackathon, 2026'
pdf=canvas.Canvas(str(OUT/'longview-three-slides.pdf'),pagesize=(W*72,H*72));pdf.setTitle('Longview | Three-slide pitch');pdf.setAuthor('Selja, Max, Jan')
font_dir=Path('/tmp/longview-browser-deps/root/usr/share/fonts/truetype/dejavu')
if not font_dir.exists():
 import reportlab
 font_dir=Path(reportlab.__file__).parent/'fonts';regular=font_dir/'Vera.ttf';bold=font_dir/'VeraBd.ttf'
else:regular=font_dir/'DejaVuSans.ttf';bold=font_dir/'DejaVuSans-Bold.ttf'
pdfmetrics.registerFont(TTFont('LV',str(regular)));pdfmetrics.registerFont(TTFont('LV-Bold',str(bold)))
class Slide:
 def __init__(self,n,dark=False):
  self.s=prs.slides.add_slide(prs.slide_layouts[6]);self.n=n;self.bg=C['dark'] if dark else C['paper'];self.dark=dark
  self.s.background.fill.solid();self.s.background.fill.fore_color.rgb=RGBColor.from_string(self.bg)
  pdf.setFillColor(HexColor('#'+self.bg));pdf.rect(0,0,W*72,H*72,fill=1,stroke=0)
 def rect(self,x,y,w,h,color,r=.0):
  sh=self.s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(C.get(color,color));sh.line.fill.background()
  if r:sh.adjustments[0]=.12
  pdf.setFillColor(HexColor('#'+C.get(color,color)));pdf.roundRect(x*72,(H-y-h)*72,w*72,h*72,r*72,fill=1,stroke=0)
 def text(self,x,y,w,h,text,size=18,color='ink',bold=False):
  tb=self.s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=tb.text_frame;tf.clear();tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0;tf.word_wrap=False
  for i,line in enumerate(text.split('\n')):
   p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='DejaVu Sans';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(C.get(color,color));p.space_after=Pt(0);p.line_spacing=1.13
  pdf.setFillColor(HexColor('#'+C.get(color,color)));pdf.setFont('LV-Bold' if bold else 'LV',size)
  for i,line in enumerate(text.split('\n')):pdf.drawString(x*72,(H-y)*72-size*.86-i*size*1.13,line)
 def line(self,x1,y1,x2,y2,color='line',width=1):
  ln=self.s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));ln.line.color.rgb=RGBColor.from_string(C.get(color,color));ln.line.width=Pt(width)
  pdf.setStrokeColor(HexColor('#'+C.get(color,color)));pdf.setLineWidth(width);pdf.line(x1*72,(H-y1)*72,x2*72,(H-y2)*72)
 def image(self,path,x,y,w,h):
  self.s.shapes.add_picture(str(path),Inches(x),Inches(y),width=Inches(w),height=Inches(h));pdf.drawImage(str(path),x*72,(H-y-h)*72,w*72,h*72)
 def footer(self,text):
  self.line(.55,7.02,12.78,7.02,'line' if not self.dark else '54725B',.5)
  self.text(.55,7.13,11.8,.2,text,6.4,'muted' if not self.dark else 'B0C5A9');self.text(12.45,7.10,.4,.2,f'{self.n:02d}',8,'muted' if not self.dark else 'B0C5A9')
 def notes(self,text):self.s.notes_slide.notes_text_frame.text=text
 def end(self):pdf.showPage()

# All shapes and text remain editable. The app screenshot is the only raster.
s=Slide(1);s.text(.6,.4,10,.3,'longview / THE PUBLIC FUNDING QUESTION',10,'muted',True)
s.text(.6,1.04,12,1.5,'What do we fund\nwhen we fund ageing?',37,'ink')
s.text(.65,3.12,5.3,.8,'A public report for people who want\nto see where research money goes.',17,'muted')
s.text(.65,4.44,5.25,.95,'One portfolio. One funding period.\nA source behind every award.\nAn estimate whose method travels with it.',14,'ink')
s.rect(6.5,1.18,6.15,5.32,'sage',.16)
s.text(6.9,1.6,5.5,.3,'THE SHARE WE WANT TO MEASURE',9,'muted',True)
s.text(6.9,2.3,5.35,1.2,'Ageing biology\n+ ageing interventions',25,'green')
s.line(6.9,3.71,12.15,3.71,'BAC7AB')
s.text(6.9,4.1,5.25,1.0,'All classified\nageing research funding',23,'ink')
s.text(6.9,5.66,5.15,.48,'Weighted by award amounts.\nCare delivery and pensions are excluded.',10,'muted')
s.footer('Method v1 · 2024 recorded awards/commitments. NIH RePORTER + Swecris, retrieved 12 Sep 2026. Selected portfolios, not national totals.')
s.notes('Track 3: communication, trust and policy; public research funding classification. Intended user: the public, with advocates sharing the report. The question is whether the biology of ageing receives a small share of research funding. We do not assume the answer. The numerator includes biology and interventions targeting ageing, with each component shown separately. The denominator is classified ageing-related research in a named government portfolio, excluding unresolved amounts from the central estimate and excluding care delivery and pensions entirely. These are commitments and awards, not annual spending. Sources: https://api.reporter.nih.gov/ and https://www.vr.se/english/swecris.html . See docs/funding/methodology.md and dataset-register.md.');s.end()

s=Slide(2,True);s.text(.6,.4,10,.3,'longview / THE WORKING REPORT',10,'B5C9AE',True)
s.text(.6,1.02,12,.8,'Two portfolios. Every award traceable.',30,'EAF0D9')
s.text(.66,2.16,4,.7,f"{r['source_count']:,}",42,'EAF0D9')
s.text(.66,2.96,4,.4,'official-source award records',13,'AEC3A5')
s.text(.66,3.62,4.0,1.0,'Sweden: 1,281\nResearch Council + Forte\n\nUnited States: 5,342\nNIA-administered awards',13,'EAF0D9')
s.text(.66,5.55,4.05,.72,'Open a portfolio. Inspect an award.\nFollow its source and funding period.',12,'AEC3A5')
from PIL import Image,ImageDraw,ImageFont
shot=Image.open(OUT/'funding-country.png');shot.save(OUT/'funding-slide-demo.png')
s.rect(5.05,2.11,7.66,4.76,'56735B',.12);s.image(OUT/'funding-slide-demo.png',5.13,2.19,7.5,4.60)
s.footer('React / TypeScript · Python / SQLite · Local objective rules · You.com + Tavily discovery · Source-linked static release; no client API keys.')
s.notes('Demo cue: start on the report, open Sweden, show the named scope and its multi-year commitment basis, inspect a source record, then show the US portfolio and its different fiscal-year basis. Search metabolism or a grant ID and follow the original award. Show the method and the visible 90% review requirement. The screenshot is the actual app. Collection covered all NIA-administered FY2024 records: 6,556 returned, with 1,214 constituent subprojects excluded, leaving 5,342. Swecris collection paged all records of VR and Forte, retaining 1,281 with FundingYear 2024. The source snapshots are complete API responses for these selected portfolios, subject to source omissions and revisions. They are not full national totals. Both providers were used for discovery crosschecks; no new paid model classification calls.');s.end()

s=Slide(3);s.text(.6,.4,10,.3,'longview / EVIDENCE BEFORE THE HEADLINE',10,'muted',True)
s.text(.6,1.02,12,.8,'The estimate waits for human review.',30,'ink')
s.rect(.6,2.04,5.73,3.75,'sage',.12);s.text(.95,2.4,5,.3,'FROZEN, BLINDED BENCHMARK',10,'muted',True)
s.text(.95,3.02,2.5,.8,'≥90%',43,'green');s.text(3.62,3.21,2.4,.5,'overall AND\nin each country',12,'muted')
s.text(.95,4.23,5,.5,f"{r['review']['overall']['reviewed']} / 60 reviewed",23,'ink')
s.text(.95,5.05,4.9,.45,'Accuracy not yet measured.\nSoftware checks do not count as human evidence.',10,'muted')
s.rect(6.58,2.04,6.13,3.75,'white',.12);s.text(6.93,2.4,5.5,.3,'WHAT CAN CHANGE THE ANSWER',10,'muted',True)
s.text(6.93,3.02,5.37,1.9,'Biology hidden under disease labels.\nMixed objectives in large awards.\nUnresolved classifications and missing amounts.\nDifferent portfolio and funding-period coverage.',13,'ink')
s.text(6.93,5.05,5.3,.48,'10 large-award checks per country,\nplus human approval of scope and arithmetic.',11,'muted')
s.rect(.6,6.03,12.11,.68,'cream',.09);s.text(.86,6.24,11.5,.35,'Next: settle the biological boundary, complete review, then test whether five readers understand the report.',10,'8A754C')
s.footer(f"Release {r['release_id']} · Pending estimates stay private. No national ranking or clinical claims. Team: Selja · Max · Jan.")
s.notes('Results: the collector, source ledger, amount-weighted aggregation, immutable report snapshots, encrypted remote review and publication gate work. Fourteen funding arithmetic and review-control tests passed during implementation; the browser verification and delivery manifest document the UI checks. These are software tests using synthetic fixtures, not a human accuracy result. The benchmark has 60 records, 30 per country, including 40 titles without ageing keywords; twelve development records are excluded. All reviews are required, with at least 54/60 and 27/30 per country. Each country also needs ten large-award checks including large exclusions, plus report-level approval. The first completed verdict is retained; changed source/label versions invalidate approvals. Remaining: human review, Andrew\'s boundary decisions, reader comprehension/adoption testing, permanent account-owned hosting and organiser submission. The call brief is in docs/funding/andrew-brief.md. Do not claim that the current estimates prove a large funding gap.');s.end()
prs.save(OUT/'longview-three-slides.pptx');pdf.save()
assert len(prs.slides)==3 and len(PdfReader(OUT/'longview-three-slides.pdf').pages)==3
(OUT/'deck-manifest.json').write_text(json.dumps({'slides':3,'funding_release':r['release_id'],'human_reviewed':r['review']['overall']['reviewed'],'human_sample':60,'editable':'All text and diagram shapes; slide 2 embeds an actual app screenshot.'},indent=2)+'\n')
# Static social preview carries no unverified numerical estimate.
im=Image.new('RGB',(1200,630),'#f7f6f0');d=ImageDraw.Draw(im)
font=lambda n:ImageFont.truetype(str(regular),n)
d.text((60,45),'LONGVIEW / AGEING RESEARCH FUNDING',font=font(18),fill='#6d8062')
for i,line in enumerate(['Where does ageing','research money go?']):d.text((60,135+i*76),line,font=font(59),fill='#193f32')
d.rounded_rectangle((60,360,1140,500),radius=12,fill='#e8eedb')
d.text((85,383),'Sweden + United States · 2024 award records',font=font(26),fill='#214c38')
d.text((85,435),'Named portfolios. Sources attached. Human review pending.',font=font(21),fill='#6b7b60')
d.text((60,556),'Recorded awards and commitments; not national spending totals.',font=font(18),fill='#7a8c6d')
(ROOT/'public/funding').mkdir(parents=True,exist_ok=True);im.save(ROOT/'public/funding/social-preview.png')
print('Created exactly 3 editable slides, a matching PDF and a static social preview.')
