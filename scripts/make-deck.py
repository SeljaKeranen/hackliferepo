"""Three editable slides and matching PDF, generated from the frozen release."""
import json,math,textwrap
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.chart import XL_CHART_TYPE,XL_LEGEND_POSITION
from pptx.chart.data import CategoryChartData
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor
from pypdf import PdfReader

ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'outputs';OUT.mkdir(exist_ok=True)
q=json.loads((ROOT/'public/release/quantitative.json').read_text());e=json.loads((ROOT/'public/release/evidence.json').read_text())
W,H=13.333,7.5
C={'paper':'F7F8F2','ink':'1B3E31','muted':'7A8C6D','line':'DCE3D2','green':'426D4D','sage':'E8EEDB','cream':'EBE4CD','white':'FFFEFA','dark':'173F34','SE':'3F8D77','US':'6F79B6','SG':'D38B4E'}
prs=Presentation();prs.slide_width=Inches(W);prs.slide_height=Inches(H)
prs.core_properties.title='Longview | Policy and population health';prs.core_properties.author='Selja, Max, Jan';prs.core_properties.subject='Stockholm AI × Longevity Hackathon, 2026'
pdf=canvas.Canvas(str(OUT/'longview-three-slides.pdf'),pagesize=(W*72,H*72));pdf.setTitle('Longview | Three-slide pitch');pdf.setAuthor('Selja, Max, Jan')
font_dir=Path('/tmp/longview-browser-deps/root/usr/share/fonts/truetype/dejavu')
if not font_dir.exists():
 import reportlab
 font_dir=Path(reportlab.__file__).parent/'fonts';regular=font_dir/'Vera.ttf';bold=font_dir/'VeraBd.ttf'
else:regular=font_dir/'DejaVuSans.ttf';bold=font_dir/'DejaVuSans-Bold.ttf'
pdfmetrics.registerFont(TTFont('LV',str(regular)));pdfmetrics.registerFont(TTFont('LV-Bold',str(bold)))
class Slide:
 def __init__(self,n,dark=False):
  self.s=prs.slides.add_slide(prs.slide_layouts[6]);self.n=n;self.bg=C['dark'] if dark else C['paper'];self.dark=dark
  self.s.background.fill.solid();self.s.background.fill.fore_color.rgb=RGBColor.from_string(self.bg)
  pdf.setFillColor(HexColor('#'+self.bg));pdf.rect(0,0,W*72,H*72,fill=1,stroke=0)
 def rect(self,x,y,w,h,color,r=.0):
  sh=self.s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(C.get(color,color));sh.line.fill.background()
  if r:sh.adjustments[0]=.12
  pdf.setFillColor(HexColor('#'+C.get(color,color)));pdf.roundRect(x*72,(H-y-h)*72,w*72,h*72,r*72,fill=1,stroke=0)
 def text(self,x,y,w,h,text,size=18,color='ink',bold=False):
  tb=self.s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=tb.text_frame;tf.clear();tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0;tf.word_wrap=False
  for i,line in enumerate(text.split('\n')):
   p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='DejaVu Sans';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(C.get(color,color));p.space_after=Pt(0);p.line_spacing=1.13
  pdf.setFillColor(HexColor('#'+C.get(color,color)));pdf.setFont('LV-Bold' if bold else 'LV',size)
  for i,line in enumerate(text.split('\n')):pdf.drawString(x*72,(H-y)*72-size*.86-i*size*1.13,line)
 def line(self,x1,y1,x2,y2,color='line',width=1):
  ln=self.s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));ln.line.color.rgb=RGBColor.from_string(C.get(color,color));ln.line.width=Pt(width)
  pdf.setStrokeColor(HexColor('#'+C.get(color,color)));pdf.setLineWidth(width);pdf.line(x1*72,(H-y1)*72,x2*72,(H-y2)*72)
 def image(self,path,x,y,w,h):
  self.s.shapes.add_picture(str(path),Inches(x),Inches(y),width=Inches(w),height=Inches(h));pdf.drawImage(str(path),x*72,(H-y-h)*72,w*72,h*72)
 def footer(self,text):
  self.line(.55,7.02,12.78,7.02,'line' if not self.dark else '54725B',.5)
  self.text(.55,7.13,11.8,.2,text,6.4,'muted' if not self.dark else 'B0C5A9');self.text(12.45,7.10,.4,.2,f'{self.n:02d}',8,'muted' if not self.dark else 'B0C5A9')
 def notes(self,text):self.s.notes_slide.notes_text_frame.text=text
 def end(self):pdf.showPage()

# Slide 1: actual observed trajectories. All text and lines are editable.
s=Slide(1);s.text(.6,.35,6,.3,'longview / THE POLICY QUESTION',10,'muted',True)
s.text(.6,.92,12,1.2,'Three countries.\nDifferent trajectories.',34,'ink',False)
s.text(.62,2.35,4.0,1.2,'A health ministry analyst needs to know\nwhat was proposed, what was funded,\nand which assumptions deserve testing.',14,'muted')
for i,c in enumerate(q['countries']):
 yy=3.95+i*.72;s.rect(.62,yy+.07,.07,.38,C[c['id']],.02);s.text(.85,yy,2.0,.25,c['name'],12,'ink',True);s.text(3.08,yy-.08,1.0,.4,f"{c['latest']:.1f}",25,c['id']);s.text(4.05,yy+.11,.5,.2,'years',8,'muted')
s.text(.85,6.22,3.5,.3,'Life expectancy at birth · 2024',9,'muted')
s.rect(5.0,2.45,7.7,4.12,'white',.12);s.text(5.35,2.75,6.8,.3,'Observed life expectancy, 1995–2024',12,'ink',True)
left,right,top,bottom=5.55,12.24,3.45,5.98
x=lambda year:left+(year-1995)/29*(right-left);y=lambda v:bottom-(v-74)/12*(bottom-top)
for val in [74,78,82,86]:s.line(left,y(val),right,y(val),'line',.55);s.text(5.15,y(val)-.07,.3,.15,str(val),7,'muted')
for year in [1995,2005,2015,2024]:s.text(x(year)-.18,6.1,.6,.2,str(year),7,'muted')
for c in q['countries']:
 hist=c['history']
 for a,b in zip(hist,hist[1:]):s.line(x(a['year']),y(a['value']),x(b['year']),y(b['value']),c['id'],2)
for i,c in enumerate(q['countries']):s.rect(5.45+i*2.3,3.2,.16,.025,c['id']);s.text(5.68+i*2.3,3.1,2,.2,c['name'],8,'muted')
s.footer('World Bank, WDI SP.DYN.LE00.IN. CC BY 4.0. Retrieved 11 Sep 2026. Period life expectancy; differences do not identify policy effects.')
s.notes('Audience: health ministry analyst. Team: Selja, Max, Jan. We compare Sweden, the United States and Singapore. The plotted values are official aggregate observations from the same WDI vintage. These differences motivate investigation; they do not show that a policy caused a longer life. The prototype joins a dated policy record to numerical context and explicit what-if assumptions. Source: https://data.worldbank.org/indicator/SP.DYN.LE00.IN . Research snapshot: 11 September 2026.');s.end()

# Slide 2: real interface, plus an editable description of the analyst's workflow.
s=Slide(2,True);s.text(.6,.35,8,.3,'longview / THE WORKING DEMO',10,'B5C9AE',True)
s.text(.6,.96,12,.7,'From a country to an analyst brief.',30,'EAF0D9')
steps=[('01','Choose a country','Observed outcomes and population\ncontext, with dates and sources.'),('02','Inspect the policy record','Separate proposals, adopted rules\nand funding commitments.'),('03','Change one assumption','Five-year forecast; ten-year scenario.\nShare the link and print the brief.')]
for i,(num,title,body) in enumerate(steps):
 yy=2.1+i*1.25;s.text(.62,yy,.5,.3,num,12,'B4C792');s.text(1.18,yy-.03,3.7,.4,title,16,'EDF2DF',True);s.text(1.18,yy+.4,3.8,.6,body,10.3,'AEC3A5')
s.rect(.6,6.0,4.1,.54,'315744',.08);s.text(.82,6.16,3.8,.2,f"{sum(v['candidates'] for v in e['countries'].values())} candidates · {len(e['findings'])} human-approved",10,'DCE7C2')
from PIL import Image
shot=Image.open(OUT/'demo-desktop.png');shot=shot.crop((0,0,1440,min(910,shot.height)));shot.save(OUT/'slide-demo-crop.png')
s.rect(5.08,1.96,7.62,4.86,'56735B',.12);s.image(OUT/'slide-demo-crop.png',5.15,2.03,7.48,4.72)
s.footer('React + TypeScript · Python + SQLite · You.com / Tavily discovery · World Bank indicators · Static browser inference; no public research endpoint.')
s.notes('Demo cue: select Singapore, open Policy evidence, show a source and the distinction between publication and effective dates once human review is complete. Until then, show the explicit review-pending state and the private team review page without marking anything approved. Return to outcomes, adjust PM2.5, switch to ten years, reset, then print the analyst brief. Policy records are not predictive inputs; sliders are user assumptions. The screenshot is the actual running app. Human verification is still pending in this release.');s.end()

# Slide 3: measured model results, honest evidence status, and the next adoption test.
s=Slide(3);s.text(.6,.35,8,.3,'longview / EVIDENCE & THE NEXT TEST',10,'muted',True)
s.text(.6,.96,11.8,.75,'Make the uncertainty part of the decision.',30,'ink')
s.rect(.6,2.0,6.0,3.7,'white',.12);s.text(.9,2.27,5.3,.35,'FIVE-YEAR HISTORICAL FORECASTS',9,'muted',True)
s.text(.9,2.8,2.6,.8,f"{q['evaluation']['test']['ridge']['mae']:.2f}",47,'green');s.text(3.08,3.03,2.5,.5,'years mean\nabsolute error',12,'muted')
for i,kind in enumerate(['ridge','no_change','trend']):
 val=q['evaluation']['test'][kind]['mae'];yy=4.05+i*.45;label={'ridge':'Regression','no_change':'No change','trend':'Previous trend'}[kind]
 s.text(.95,yy,1.45,.22,label,9,'muted');s.rect(2.43,yy+.035,val/1.4*2.85,.15,'green' if kind=='ridge' else 'C5D0B7',.03);s.text(5.47,yy-.01,.7,.3,f'{val:.2f}',11,'ink')
s.text(.9,5.4,5.2,.2,f"{q['evaluation']['test']['ridge']['n']:,} tests · {q['evaluation']['training_countries']} training countries / economies",8,'muted')
s.rect(6.86,2.0,5.86,3.7,'sage',.12);s.text(7.17,2.27,5.2,.3,'HUMAN EVIDENCE VERIFICATION',9,'muted',True)
s.text(7.17,2.86,2.55,.8,'≥90%',43,'green');s.text(9.72,3.01,2.55,.5,'required initial\nsample pass rate',11,'muted')
rate='Not measured' if e['benchmark']['accuracy'] is None else f"{e['benchmark']['accuracy']*100:.1f}%"
s.text(7.17,4.03,5.1,.5,f"{e['benchmark']['reviewed']} / {e['benchmark']['sample_size']} reviewed · {rate}",15,'ink',True)
s.text(7.17,4.65,5.1,.65,'Source, date, classification and claim.\nThe initial sample is frozen before correction.\nUnapproved claims stay out of the public release.',10,'muted')
s.rect(.6,5.95,12.12,.72,'cream',.09)
s.text(.85,6.12,11.6,.45,f"Known failure: US error {q['evaluation']['country_holdout']['USA']['ridge']['mae']:.2f} years vs {q['evaluation']['country_holdout']['USA']['no_change']['mae']:.2f} for no change. Wider interval coverage: {q['evaluation']['test']['ridge']['coverage']*100:.1f}% (90% target).",10,'8A754C')
s.text(.65,6.77,12,.18,'Next pilot: test analyst interpretation and review comparable fundamental-ageing research funding. Team: Selja · Max · Jan.',8.0,'muted')
s.footer(f"Measured from release {q['release_id']} / {e['release_id']}, 11 Sep 2026. Current revised data; non-causal model. No clinical claims.")
s.notes('The ridge model beats both baselines on the wider historical test, using model selection fixed on a pre-test validation period. Sweden, the US and Singapore were withheld from every fitting, tuning and calibration set. The US regression fails against both baselines. The displayed interval is an empirical error band: nominal 90% coverage achieved only 74.9% in the wider test. There are six overlapping historical origins per held-out country, so these are small dependent samples. The evidence requirement is different: at least 90% of the frozen initial findings must pass four human checks. At generation time no human reviews have been imported. Two supplemental funding leads are outside the initial 34-record benchmark and also require approval. The proposed adoption pilot measures source-checking time and whether analysts correctly understand conditional scenarios; no adoption benefit has been measured yet. Sources: public/release/quantitative.json; public/release/evidence.json; docs/model-evaluation.json.');s.end()
prs.save(OUT/'longview-three-slides.pptx');pdf.save()
assert len(prs.slides)==3 and len(PdfReader(OUT/'longview-three-slides.pdf').pages)==3
(OUT/'deck-manifest.json').write_text(json.dumps({'slides':3,'quantitative_release':q['release_id'],'evidence_release':e['release_id'],'human_reviewed':e['benchmark']['reviewed'],'human_sample':e['benchmark']['sample_size'],'editable':'All text, boxes and data lines; slide 2 embeds a real app screenshot.'},indent=2)+'\n')
print('Created exactly 3 slides: outputs/longview-three-slides.pptx and .pdf')
